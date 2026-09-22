"""共通の描画部品（KPIカード・表・グラフ）。python-pptxのみに依存し、DBは知らない。"""

from __future__ import annotations

from datetime import date, datetime

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.presentation import Presentation
from pptx.slide import Slide
from pptx.util import Pt

DASH = "―"

COLOR_POSITIVE = RGBColor(0x1E, 0x7E, 0x34)  # 緑：達成・プラス
COLOR_NEGATIVE = RGBColor(0xC0, 0x39, 0x2B)  # 赤：未達・マイナス
COLOR_NEUTRAL = RGBColor(0x33, 0x33, 0x33)

LAYOUT_TITLE = 0  # 表紙用（タイトル＋サブタイトル）
LAYOUT_CONTENT = 5  # 通常スライド用（タイトルのみ、本文は自作）


def fmt_millions(amount: float | int | None) -> str:
    """金額を百万円単位・小数点1桁の文字列にする。"""
    if amount is None:
        return DASH
    return f"{amount / 1_000_000:.1f}"


def fmt_percent(ratio: float | None, *, signed: bool = False) -> str:
    """比率を小数点1桁の％表記にする。"""
    if ratio is None:
        return DASH
    pct = ratio * 100
    return f"{pct:+.1f}%" if signed else f"{pct:.1f}%"


def ratio_color(ratio: float | None) -> RGBColor:
    """前月比・前年比などの正負に応じた色（プラス=緑、マイナス=赤）。"""
    if ratio is None:
        return COLOR_NEUTRAL
    return COLOR_POSITIVE if ratio >= 0 else COLOR_NEGATIVE


def achievement_color(ratio: float | None) -> RGBColor:
    """予算達成率に応じた色（達成=緑、未達=赤）。"""
    if ratio is None:
        return COLOR_NEUTRAL
    return COLOR_POSITIVE if ratio >= 1.0 else COLOR_NEGATIVE


def new_slide(prs: Presentation, layout_index: int, title: str | None = None) -> Slide:
    layout = prs.slide_layouts[layout_index]
    slide = prs.slides.add_slide(layout)
    if title is not None and slide.shapes.title is not None:
        slide.shapes.title.text = title
    return slide


def set_notes(slide: Slide, *, month: date, region: str | None, generated_at: datetime) -> None:
    """ノート欄に集計条件（対象月・フィルタ・生成日時）を記載する。"""
    lines = [
        f"対象月: {month:%Y-%m}",
        f"拠点: {region or '全体'}",
        f"生成日時: {generated_at:%Y-%m-%d %H:%M}",
    ]
    slide.notes_slide.notes_text_frame.text = "\n".join(lines)


def add_kpi_card(
    slide: Slide,
    left,
    top,
    width,
    height,
    label: str,
    value: str,
    *,
    value_color: RGBColor | None = None,
    sub: str | None = None,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    r0.text = label
    r0.font.size = Pt(12)
    r0.font.color.rgb = COLOR_NEUTRAL

    p1 = tf.add_paragraph()
    r1 = p1.add_run()
    r1.text = value
    r1.font.size = Pt(28)
    r1.font.bold = True
    r1.font.color.rgb = value_color or COLOR_NEUTRAL

    if sub:
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = sub
        r2.font.size = Pt(11)
        r2.font.color.rgb = COLOR_NEUTRAL

    return box


def add_bullets(slide: Slide, left, top, width, height, lines: list[str], *, font_size: int = 16):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True

    items = lines or ["特筆すべき事項はありません。"]
    for i, line in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"・{line}"
        for r in p.runs:
            r.font.size = Pt(font_size)

    return box


def add_table(
    slide: Slide,
    left,
    top,
    width,
    height,
    headers: list[str],
    rows: list[list[str]],
    *,
    font_color_fn=None,
):
    """headers/rowsは表示用の文字列を渡す。font_color_fn(row_idx, col_idx, raw_row) -> RGBColor|None"""
    n_rows = len(rows) + 1
    n_cols = len(headers)
    gframe = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = gframe.table

    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = str(h)
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(12)

    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(ri + 1, ci)
            cell.text = DASH if val is None else str(val)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.RIGHT
                for r in p.runs:
                    r.font.size = Pt(11)
                    if font_color_fn:
                        color = font_color_fn(ri, ci, row)
                        if color:
                            r.font.color.rgb = color

    return table


def add_line_chart(slide: Slide, left, top, width, height, categories: list[str], series: dict[str, list[float]]):
    chart_data = CategoryChartData()
    chart_data.categories = categories
    for name, values in series.items():
        chart_data.add_series(name, values)

    gframe = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, left, top, width, height, chart_data)
    chart = gframe.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    return chart


def add_bar_chart(
    slide: Slide,
    left,
    top,
    width,
    height,
    categories: list[str],
    series: dict[str, list[float]],
    *,
    horizontal: bool = True,
):
    chart_data = CategoryChartData()
    chart_data.categories = categories
    for name, values in series.items():
        chart_data.add_series(name, values)

    chart_type = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gframe = slide.shapes.add_chart(chart_type, left, top, width, height, chart_data)
    chart = gframe.chart
    chart.has_legend = len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    return chart


def add_pie_chart(slide: Slide, left, top, width, height, categories: list[str], values: list[float]):
    chart_data = CategoryChartData()
    chart_data.categories = categories
    chart_data.add_series("売上", values)

    gframe = slide.shapes.add_chart(XL_CHART_TYPE.PIE, left, top, width, height, chart_data)
    chart = gframe.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.RIGHT
    chart.legend.include_in_layout = False
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.show_percentage = True
    plot.data_labels.show_value = False
    return chart
