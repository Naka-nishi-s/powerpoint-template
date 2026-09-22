"""1 表紙：タイトル、対象月、作成日"""

from __future__ import annotations

from pptx.presentation import Presentation

from .. import builder
from ..metrics import ReportData


def build_cover(prs: Presentation, data: ReportData) -> None:
    slide = builder.new_slide(prs, builder.LAYOUT_TITLE)
    slide.shapes.title.text = "月次売上レポート"

    subtitle = slide.placeholders[1]
    lines = [f"対象月: {data.month:%Y年%m月}"]
    if data.region:
        lines.append(f"拠点: {data.region}")
    lines.append(f"作成日: {data.generated_at:%Y-%m-%d}")
    subtitle.text_frame.text = "\n".join(lines)

    builder.set_notes(slide, month=data.month, region=data.region, generated_at=data.generated_at)
