"""6 商品ランキング：売上トップ10と前年比ワースト5（表を2つ並べる）"""

from __future__ import annotations

from pptx.presentation import Presentation
from pptx.util import Inches, Pt

from .. import builder
from ..metrics import ReportData


def build_product(prs: Presentation, data: ReportData) -> None:
    slide = builder.new_slide(prs, builder.LAYOUT_CONTENT, title="商品ランキング")

    top10 = data.product.top10
    worst5 = data.product.worst5

    label = slide.shapes.add_textbox(Inches(0.5), Inches(1.35), Inches(6.0), Inches(0.35))
    label.text_frame.text = "売上トップ10"
    label.text_frame.paragraphs[0].runs[0].font.bold = True
    label.text_frame.paragraphs[0].runs[0].font.size = Pt(14)

    headers1 = ["順位", "商品名", "カテゴリ", "売上（百万円）", "数量"]
    rows1 = []
    for i, row in enumerate(top10.itertuples(), start=1):
        rows1.append([str(i), row.product_name, row.category, builder.fmt_millions(row.amount), str(int(row.quantity))])
    builder.add_table(slide, Inches(0.5), Inches(1.75), Inches(6.0), Inches(5.0), headers1, rows1)

    label2 = slide.shapes.add_textbox(Inches(6.8), Inches(1.35), Inches(6.0), Inches(0.35))
    label2.text_frame.text = "前年比ワースト5"
    label2.text_frame.paragraphs[0].runs[0].font.bold = True
    label2.text_frame.paragraphs[0].runs[0].font.size = Pt(14)

    headers2 = ["順位", "商品名", "カテゴリ", "前年比"]
    rows2 = []
    for i, row in enumerate(worst5.itertuples(), start=1):
        rows2.append([str(i), row.product_name, row.category, builder.fmt_percent(row.yoy_ratio, signed=True)])

    def color_fn(row_idx, col_idx, raw_row):
        if col_idx == 3:
            return builder.ratio_color(worst5.iloc[row_idx]["yoy_ratio"])
        return None

    builder.add_table(
        slide, Inches(6.8), Inches(1.75), Inches(6.0), Inches(5.0), headers2, rows2, font_color_fn=color_fn
    )

    builder.set_notes(slide, month=data.month, region=data.region, generated_at=data.generated_at)
