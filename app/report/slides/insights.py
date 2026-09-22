"""8 所見・課題：自動生成した所見と、担当者が記入する欄"""

from __future__ import annotations

from pptx.enum.shapes import MSO_SHAPE
from pptx.presentation import Presentation
from pptx.util import Inches, Pt

from .. import builder
from ..insights import generate_insights
from ..metrics import ReportData


def build_insights(prs: Presentation, data: ReportData) -> None:
    slide = builder.new_slide(prs, builder.LAYOUT_CONTENT, title="所見・課題")

    texts = generate_insights(data) if data.with_insights else []
    builder.add_bullets(slide, Inches(0.5), Inches(1.5), Inches(12.3), Inches(2.6), texts, font_size=16)

    label = slide.shapes.add_textbox(Inches(0.5), Inches(4.3), Inches(4.0), Inches(0.35))
    label.text_frame.text = "メモ欄（担当者記入）"
    label.text_frame.paragraphs[0].runs[0].font.bold = True
    label.text_frame.paragraphs[0].runs[0].font.size = Pt(14)

    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(4.7), Inches(12.3), Inches(2.3))
    box.fill.background()
    box.line.color.rgb = builder.COLOR_NEUTRAL
    box.line.width = Pt(0.75)
    box.text_frame.word_wrap = True
    box.text_frame.text = ""

    builder.set_notes(slide, month=data.month, region=data.region, generated_at=data.generated_at)
