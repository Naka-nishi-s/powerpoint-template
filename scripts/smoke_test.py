"""DB接続なしでreportパッケージの生成ロジックを検証するスモークテスト。
queries.pyの各fetch関数をモックし、generate_slide/generate_reportが
例外なくpptxバイト列を返すことを確認する。
"""

import sys
from datetime import date
from pathlib import Path
from unittest import mock

import pandas as pd

APP_DIR = Path(__file__).resolve().parent.parent / "app"
sys.path.insert(0, str(APP_DIR))

MONTHS = [date(2024, 9, 1) if i == 0 else None for i in range(24)]
def month_add(d, n):
    total = d.year * 12 + (d.month - 1) + n
    y, m = divmod(total, 12)
    return date(y, m + 1, 1)

ALL_MONTHS = [month_add(date(2024, 9, 1), i) for i in range(24)]
STORES = [
    (1, "札幌店", "北海道", "店舗"),
    (2, "仙台店", "東北", "店舗"),
    (3, "東京本店", "関東", "店舗"),
    (10, "オンラインストア", "関東", "EC"),
]
CATEGORIES = ["食品・飲料", "日用雑貨", "家電", "アパレル", "インテリア"]
PRODUCTS = [(i, f"商品{i:03d}", CATEGORIES[i % 5]) for i in range(1, 11)]


def fake_company_monthly(engine, months, region):
    rows = []
    for m in months:
        if m not in ALL_MONTHS:
            continue
        base = 50_000_000 + ALL_MONTHS.index(m) * 300_000
        rows.append({"month": m, "amount": base, "gross_profit": base * 0.3, "quantity": 5000, "order_count": 4000})
    return pd.DataFrame(rows, columns=["month", "amount", "gross_profit", "quantity", "order_count"])


def fake_category_monthly(engine, months, region):
    rows = []
    for m in months:
        if m not in ALL_MONTHS:
            continue
        for i, cat in enumerate(CATEGORIES):
            amt = 8_000_000 + i * 500_000 + ALL_MONTHS.index(m) * 10_000
            rows.append({"month": m, "category": cat, "amount": amt, "gross_profit": amt * 0.3})
    return pd.DataFrame(rows, columns=["month", "category", "amount", "gross_profit"])


def fake_store_monthly(engine, months, region):
    rows = []
    for m in months:
        if m not in ALL_MONTHS:
            continue
        for sid, name, reg, ch in STORES:
            if region and reg != region:
                continue
            amt = 5_000_000 + sid * 100_000
            target = 5_500_000
            rows.append({"month": m, "store_id": sid, "store_name": name, "region": reg, "channel": ch, "amount": amt, "target_amount": target})
    return pd.DataFrame(rows, columns=["month", "store_id", "store_name", "region", "channel", "amount", "target_amount"])


def fake_product_monthly(engine, months, region):
    rows = []
    for m in months:
        if m not in ALL_MONTHS:
            continue
        for pid, name, cat in PRODUCTS:
            amt = 1_000_000 + pid * 50_000 + (ALL_MONTHS.index(m) - 12) * 5_000
            rows.append({"month": m, "product_id": pid, "product_name": name, "category": cat, "amount": max(amt, 10000), "quantity": 100 + pid})
    return pd.DataFrame(rows, columns=["month", "product_id", "product_name", "category", "amount", "quantity"])


def fake_customer_monthly(engine, months, region):
    rows = []
    for m in months:
        if m not in ALL_MONTHS:
            continue
        for cid in range(1, 21):
            rows.append({
                "month": m,
                "customer_id": cid,
                "customer_type": "法人" if cid % 4 == 0 else "個人",
                "is_new": cid == 1,
                "amount": 200_000 + cid * 1000,
            })
    return pd.DataFrame(rows, columns=["month", "customer_id", "customer_type", "is_new", "amount"])


def fake_available_months(engine):
    return ALL_MONTHS


def fake_regions(engine):
    return sorted({reg for _, _, reg, _ in STORES})


def run():
    with mock.patch("report.queries.fetch_company_monthly", fake_company_monthly), \
         mock.patch("report.queries.fetch_category_monthly", fake_category_monthly), \
         mock.patch("report.queries.fetch_store_monthly", fake_store_monthly), \
         mock.patch("report.queries.fetch_product_monthly", fake_product_monthly), \
         mock.patch("report.queries.fetch_customer_monthly", fake_customer_monthly), \
         mock.patch("report.queries.fetch_available_months", fake_available_months), \
         mock.patch("report.queries.fetch_regions", fake_regions):

        from report.generate import generate_report, generate_slide
        from report.slides import SLIDES
        from pptx import Presentation
        import io

        target_month = date(2026, 8, 1)

        for slide_id in SLIDES:
            content = generate_slide(target_month, slide_id, None, engine=object())
            prs = Presentation(io.BytesIO(content))
            assert len(prs.slides) == 1, f"{slide_id}: expected 1 slide, got {len(prs.slides)}"
            print(f"OK generate_slide[{slide_id}] -> {len(content)} bytes, 1 slide")

        content = generate_report(target_month, None, None, True, engine=object())
        prs = Presentation(io.BytesIO(content))
        assert len(prs.slides) == 8, f"expected 8 slides, got {len(prs.slides)}"
        print(f"OK generate_report(all) -> {len(content)} bytes, {len(prs.slides)} slides")

        content = generate_report(target_month, "関東", ["trend", "store"], False, engine=object())
        prs = Presentation(io.BytesIO(content))
        assert len(prs.slides) == 4, f"expected 4 slides (cover+summary+trend+store), got {len(prs.slides)}"
        print(f"OK generate_report(subset, region) -> {len(content)} bytes, {len(prs.slides)} slides")

        print("ALL SMOKE TESTS PASSED")


if __name__ == "__main__":
    run()
